# -*- coding: utf-8 -*-
"""Python - Orienté objet (enrichi) — génération du deck au style bstorm Training."""
from pathlib import Path

from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

import deck_lib as L
from deck_lib import (ASSETS, GREEN, PURPLE, add_bonus_tag, add_code,
                      add_footer, add_paragraphs, add_slide_title, add_table,
                      add_textbox, blank_slide, new_presentation,
                      rich_runs, set_bullet, style_run)

OUT = Path(__file__).parent.parent / "Python - Orienté objet (enrichi).pptx"


# ================================================================ types de slides
def slide_titre(prs):
    s = blank_slide(prs)
    box, tf = add_textbox(s, 0, 128, 720, 70)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    style_run(p.add_run(), "Python", font=L.TITLE_FONT, size=52, color=GREEN, bold=True)
    box, tf = add_textbox(s, 0, 200, 720, 40)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    style_run(p.add_run(), "Orienté objet", font=L.TITLE_FONT, size=27, color=PURPLE, bold=True)
    s.shapes.add_picture(str(ASSETS / "logo_training.png"), Pt(14), Pt(322), height=Pt(72))


def slide_marque(prs):
    s = blank_slide(prs)
    s.shapes.add_picture(str(ASSETS / "qui_sommes_nous.png"), 0, 0,
                         width=L.SLIDE_W, height=L.SLIDE_H)


def slide_section(prs, num, titre):
    s = blank_slide(prs)
    box, tf = add_textbox(s, 0, 178, 720, 50)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    style_run(p.add_run(), titre, font=L.TITLE_FONT, size=28, color=PURPLE, bold=False)
    add_footer(s, num)


def slide_toc(prs, num, titre, col_gauche, col_droite):
    s = blank_slide(prs)
    add_slide_title(s, titre)
    for items, x in ((col_gauche, 42), (col_droite, 375)):
        box, tf = add_textbox(s, x, 92, 320, 280)
        for i, (lvl, text) in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = 1.0
            if lvl == 0:
                set_bullet(p, "➔", 0, size_pct=85)
                p.space_before = Pt(9)
                p.space_after = Pt(3)
                rich_runs(p, text, size=14.5)
            else:
                set_bullet(p, "◆", 1, size_pct=70)
                p.space_after = Pt(2.5)
                rich_runs(p, text, size=13)
    add_footer(s, num)


def slide_contenu(prs, num, titre, blocks, bonus=False, body_size=15):
    s = blank_slide(prs)
    add_slide_title(s, titre)
    if bonus:
        add_bonus_tag(s)
    y = 86.0
    for block in blocks:
        kind = block[0]
        if kind == "txt":
            h = add_paragraphs(s, block[1], 30, y, 660, size=body_size)
            y += h + 8
        elif kind == "code":
            code = block[1]
            opts = block[2] if len(block) > 2 else {}
            x = opts.get("x", 60)
            w = opts.get("w", 600)
            h = add_code(s, code, x, y, w, size=opts.get("size", 12))
            y += h + 10
        elif kind == "table":
            headers, rows, widths = block[1], block[2], block[3]
            opts = block[4] if len(block) > 4 else {}
            h = add_table(s, headers, rows, widths, opts.get("x", 30), y,
                          size=opts.get("size", 12), row_h=opts.get("row_h", 21),
                          mono_cols=opts.get("mono_cols", ()))
            y += h + 10
        elif kind == "gap":
            y += block[1]
        elif kind == "note":
            box, tf = add_textbox(s, 30, y, 660, 20)
            p = tf.paragraphs[0]
            rich_runs(p, block[1], size=12.5, color=L.GRAY)
            y += 20
    add_footer(s, num)
    return y


def slide_merci(prs, num):
    s = blank_slide(prs)
    box, tf = add_textbox(s, 0, 178, 720, 50)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    style_run(p.add_run(), "Merci pour votre attention.", font=L.TITLE_FONT,
              size=28, color=PURPLE)
    add_footer(s, num)


# ================================================================ le contenu
def build():
    prs = new_presentation()
    n = 0

    def C(titre, *blocks, bonus=False, body_size=15):
        nonlocal n
        n += 1
        slide_contenu(prs, n, titre, list(blocks), bonus=bonus, body_size=body_size)

    def SEC(titre):
        nonlocal n
        n += 1
        slide_section(prs, n, titre)

    # ---- titre + marque ----------------------------------------------------
    n += 1
    slide_titre(prs)
    n += 1
    slide_marque(prs)

    # ---- table des matières ------------------------------------------------
    n += 1
    slide_toc(prs, n, "Table des matières",
              [(0, "Introduction"),
               (1, "paradigme orienté objet"),
               (1, "coder dans plusieurs fichiers"),
               (0, "Les classes"),
               (1, "self, __init__"),
               (1, "attribut d'instance vs de classe"),
               (1, "pas de surcharge de méthodes"),
               (0, "Encapsulation"),
               (1, "visibilité _ et __"),
               (1, "propriétés et validation"),
               (0, "Data model"),
               (1, "__str__ et __repr__"),
               (1, "__eq__, opérateurs, conteneurs")],
              [(0, "Héritage"),
               (1, "super(), redéfinition"),
               (1, "héritage multiple et MRO"),
               (1, "composition vs héritage"),
               (1, "exceptions personnalisées"),
               (0, "Les classes abstraites"),
               (0, "Les membres statiques"),
               (1, "@staticmethod, @classmethod"),
               (0, "Les énumérations"),
               (0, "Les interfaces"),
               (1, "duck typing, ABC, Protocol")])

    # ======================================================== INTRODUCTION
    SEC("Introduction")

    C("Paradigme orienté objet",
      ("txt", [("p", "Le paradigme orienté objet a pour objectif de définir des briques logicielles appelées des **objets** — représentant des entités physiques, des concepts ou des idées."),
               ("p", "Chaque objet possède :"),
               ("b0", "une **structure interne** — son état : des données, les *attributs*"),
               ("b0", "un **comportement** — ce qu'il sait faire : des *méthodes*"),
               ("p", "L'intérêt : regrouper au même endroit les données et le code qui les manipule — le programme se structure autour du **métier** (une voiture, un client, une tâche) plutôt qu'autour de fonctions éparpillées.")]))

    C("Le concept d'objet",
      ("txt", [("p", "Un objet est une entité modélisant une idée, un concept ou une entité physique du monde réel. Il est caractérisé par :"),
               ("b0", "un **état** — la couleur, l'année et le kilométrage d'une voiture"),
               ("b0", "un **comportement** — rouler, klaxonner, freiner")]),
      ("table", ["Monde réel", "État (attributs)", "Comportement (méthodes)"],
       [["une voiture", "marque, annee, couleur", "rouler(), klaxonner()"],
        ["un compte en banque", "titulaire, solde", "deposer(), retirer()"],
        ["une tâche (projet Todo)", "titre, faite", "terminer(), afficher()"]],
       [170, 240, 250], {"mono_cols": (1, 2), "row_h": 23}))

    C("Coder dans plusieurs fichiers",
      ("txt", [("p", "Pour développer de manière structurée, nous allons coder dans plusieurs fichiers — un import les relie :"),
               ("b0", "`from` — le fichier qui contient le code à importer (répertoire en package + nom du fichier)"),
               ("b0", "`import` — l'élément à rendre disponible : une fonction, une variable ou une classe")]),
      ("code", """
# fichier : models/voiture.py
class Voiture:
    def __init__(self, marque):
        self.marque = marque

# fichier : principal.py
from models.voiture import Voiture     # package.fichier  import  classe

v = Voiture("Renault")
"""))

    C("Limiter l'exécution du code",
      ("txt", [("p", "À l'import d'un fichier, **tout son code de niveau principal s'exécute**. Pour réserver une partie (tests, démonstration) au lancement direct du fichier, on utilise la variable `__name__` :"),
               ("b0", "lancé directement → `__name__` vaut `\"__main__\"`"),
               ("b0", "importé → `__name__` vaut le nom du module : le bloc est ignoré")]),
      ("code", """
# fichier : outils.py
def aide():
    print("Je suis l'outil")

if __name__ == "__main__":      # vrai SEULEMENT si on lance CE fichier
    print("Zone de test du module...")
    aide()
"""))

    # ======================================================== LES CLASSES
    SEC("Les classes")

    C("Définition d'une classe",
      ("txt", [("p", "Une classe définit un **type d'objet** : ses états (attributs) et ses comportements (méthodes)."),
               ("b0", "La classe est un **plan de construction** (un moule) — elle ne contient pas de données concrètes"),
               ("b0", "Un objet créé à partir d'une classe est une **instance** de cette classe"),
               ("b0", "On peut créer autant d'objets que l'on désire avec une même classe")]),
      ("code", """
class Voiture:                 # le mot clef « class » + PascalCase
    pass

v1 = Voiture()                 # une instance...
v2 = Voiture()                 # ... une autre, indépendante
print(type(v1))                # <class '__main__.Voiture'>
"""))

    C("« self » — qui parle ?",
      ("txt", [("p", "Le premier paramètre de toute méthode reçoit **l'instance sur laquelle on appelle la méthode**. Par convention, il s'appelle `self`."),
               ("p", "La notation pointée n'est qu'un raccourci : Python passe l'objet **avant le point** comme premier argument.")]),
      ("code", """
class Voiture:
    def klaxonner(self):
        print("Tut tut !")

v = Voiture()
v.klaxonner()                  # ce que Python exécute réellement :
Voiture.klaxonner(v)           # ... la même chose — self reçoit v !
"""),
      ("txt", [("p", "L'oubli de `self` dans la signature est l'erreur n°1 du chapitre : `TypeError: klaxonner() takes 0 positional arguments but 1 was given`.")]))

    C("« __init__ » — initialiser l'objet",
      ("txt", [("p", "La méthode spéciale `__init__` est appelée **automatiquement à la création** de chaque objet. C'est elle qui crée les **attributs d'instance** — les données propres à chaque objet."),
               ("p", "Ses paramètres reçoivent les valeurs passées à l'instanciation.")]),
      ("code", """
class Voiture:
    def __init__(self, marque, annee):
        self.marque = marque       # attribut d'INSTANCE : propre à chaque objet
        self.annee = annee

v1 = Voiture("Renault", 2019)      # __init__ est appelée ici, self = v1
v2 = Voiture("Kia", 2024)
print(v1.marque, v2.marque)        # Renault Kia — chacun les siens
"""))

    C("Attribut d'instance VS attribut de classe",
      ("txt", [("b0", "**Attribut d'instance** — créé dans `__init__` via `self.` : chaque objet a **sa** valeur"),
               ("b0", "**Attribut de classe** — déclaré directement dans la classe : **une seule** valeur, partagée par toutes les instances")]),
      ("code", """
class Voiture:
    nb_roues = 4                   # attribut de CLASSE : partagé par toutes

    def __init__(self, marque):
        self.marque = marque       # attribut d'INSTANCE : propre à chacune

v1 = Voiture("Renault")
v2 = Voiture("Kia")
print(v1.nb_roues, v2.nb_roues)    # 4 4 — la même valeur, stockée UNE fois
print(v1.marque, v2.marque)        # Renault Kia
"""))

    C("Le piège du masquage",
      ("txt", [("p", "Affecter via l'instance **ne modifie pas** l'attribut de classe : cela crée un attribut d'instance qui le **masque** pour cet objet-là seulement.")]),
      ("code", """
class Compteur:
    total = 0                    # attribut de classe

c = Compteur()
c.total = 10                     # crée un attribut d'INSTANCE qui MASQUE total
print(Compteur.total)            # 0  — l'attribut de classe n'a pas bougé
print(c.total)                   # 10 — cette instance voit le sien

Compteur.total = 99              # modifier la valeur partagée : via la CLASSE
"""),
      ("txt", [("p", "Règle de lecture : Python cherche d'abord dans l'instance, puis dans la classe. Règle d'écriture : `instance.x = ...` écrit **toujours** dans l'instance.")]))

    C("Les méthodes",
      ("txt", [("p", "Les méthodes sont les procédures et fonctions définies au sein de la classe : elles portent les **comportements**. Le premier paramètre, obligatoire, est `self` — il donne accès aux attributs de l'objet."),
               ("p", "Convention de nommage : le « snake_case », comme les fonctions.")]),
      ("code", """
class Voiture:
    def __init__(self, marque, km):
        self.marque = marque
        self.km = km

    def rouler(self, distance):
        self.km += distance            # la méthode lit ET modifie SON objet
        print(f"{self.marque} : {self.km} km au compteur")

v = Voiture("Renault", 68000)
v.rouler(150)                          # Renault : 68150 km au compteur
"""))

    C("Pas de surcharge de méthodes !",
      ("txt", [("p", "Contrairement à C# et Java, Python n'a **pas de surcharge** : redéfinir une méthode avec d'autres paramètres... **écrase la précédente en silence**. Le remplacement : les valeurs par défaut (et `*args`).")]),
      ("code", """
class Voiture:
    def rouler(self):
        return "je roule"
    def rouler(self, km):               # écrase la 1re en silence !
        return f"je roule {km} km"

# Voiture().rouler()        # TypeError — il ne reste QUE la version à km
print(Voiture().rouler(10)) # je roule 10 km

class Voiture2:
    def rouler(self, km=None):          # LA solution : paramètre optionnel
        return "je roule" if km is None else f"je roule {km} km"

print(Voiture2().rouler(), "|", Voiture2().rouler(5))   # je roule | je roule 5 km
""", {"size": 11}))

    # ======================================================== ENCAPSULATION
    SEC("Encapsulation")

    C("Le but de l'encapsulation",
      ("txt", [("p", "L'encapsulation est un concept primordial de l'orienté objet : **cacher les données internes** d'un objet, et n'y donner accès qu'à travers des méthodes contrôlées."),
               ("b0", "elle permet un **couplage faible** entre les classes"),
               ("b0", "elle garantit l'**intégrité** des données internes (impossible d'y mettre n'importe quoi)"),
               ("p", "Image : le tableau de bord expose un compteur et des pédales — personne ne touche directement aux pistons.")]))

    C("La visibilité : _ et __",
      ("txt", [("p", "Python n'a **pas** de modificateurs de visibilité (`private`, `protected`...). Il utilise des **conventions d'écriture** :")]),
      ("table", ["Préfixe", "Signification", "Réalité technique"],
       [["_attribut", "« privé par convention » — n'y touchez pas", "accessible normalement ; l'IDE avertit"],
        ["__attribut", "« fortement privé »", "renommé en _Classe__attribut (name mangling) — contournable, mais le message est clair"]],
       [110, 230, 320], {"mono_cols": (0,), "row_h": 30}),
      ("gap", 4),
      ("txt", [("p", "À retenir : ce sont des **conventions de politesse outillées**, pas une sécurité. L'encapsulation en Python repose sur la discipline... et sur les propriétés.")]))

    C("Les propriétés",
      ("txt", [("p", "Une propriété est un ensemble de méthodes qui **contrôlent l'accès** aux attributs de la classe :"),
               ("b0", "l'**accesseur** (getter) — consulter la valeur"),
               ("b0", "le **mutateur** (setter) — modifier la valeur, avec contrôle"),
               ("p", "Vu de l'extérieur, on continue d'écrire `objet.attribut` — la propriété s'interpose discrètement : c'est toute l'élégance du mécanisme.")]))

    C("La syntaxe : @property et @setter",
      ("code", """
class Car:
    def __init__(self, name, hp):
        self.name = name
        self.hp = hp                   # passe DÉJÀ par le setter ci-dessous !

    @property
    def hp(self):                      # l'accesseur : objet.hp en LECTURE
        return self.__hp

    @hp.setter
    def hp(self, value):               # le mutateur : objet.hp = ... en ÉCRITURE
        if value <= 0:
            raise ValueError(f"Puissance invalide : {value}")
        self.__hp = value

c = Car("Clio", 90)
print(c.hp)            # 90 — on dirait un attribut, c'est une méthode déguisée
"""))

    C("La règle d'or du setter",
      ("txt", [("p", "⚠ Le piège du setter « silencieux » : un `if` **sans else** qui refuse la valeur... ne fait rien. L'attribut n'est **jamais créé**, et le programme plante plus tard, loin du vrai bug.")]),
      ("code", """
# ❌ Refus SILENCIEUX
@hp.setter
def hp(self, value):
    if value > 0:
        self.__hp = value      # Car("Clio", -5) → __hp jamais créé...
                               # ... et AttributeError à la 1re lecture, ailleurs !

# ✅ Un setter qui refuse LÈVE une erreur — immédiate, claire, localisée
@hp.setter
def hp(self, value):
    if value <= 0:
        raise ValueError(f"Puissance invalide : {value}")
    self.__hp = value
""", {"size": 11.5}),
      ("txt", [("p", "Règle : **valider, c'est refuser bruyamment** — `raise ValueError(...)`, jamais un échec muet.")]))

    C("Propriété calculée (lecture seule)", bonus=True, *[
      ("txt", [("p", "Une propriété sans setter expose une valeur **calculée** à la demande — en lecture seule :")]),
      ("code", """
class Car:
    def __init__(self, name, prix_ht):
        self.name = name
        self.prix_ht = prix_ht

    @property
    def prix_ttc(self):                    # calculée à chaque lecture
        return round(self.prix_ht * 1.21, 2)

c = Car("Clio", 10000)
print(c.prix_ttc)          # 12100.0
# c.prix_ttc = 5           # ❌ AttributeError : pas de setter → lecture seule
""")])

    # ======================================================== DATA MODEL
    SEC("Data model")

    C("Les attributs spéciaux",
      ("txt", [("p", "Les attributs spéciaux (entourés de doubles underscores) sont fournis par Python, en lecture seule :")]),
      ("table", ["Nom", "Accessible sur", "Renvoie"],
       [["__name__", "classe", "le nom de la classe"],
        ["__bases__", "classe", "le tuple des classes parentes"],
        ["__class__", "objet", "la classe qui a créé l'objet"],
        ["__dict__", "objet", "le dictionnaire des attributs **de l'instance**"],
        ["__doc__", "classe et objet", "la docstring de la classe"]],
       [110, 140, 410], {"mono_cols": (0,), "row_h": 22}),
      ("gap", 4),
      ("txt", [("p", "`v.__dict__` est un excellent outil de débogage : il montre exactement ce que **cet objet** stocke — et lui seul (pas les méthodes, pas les attributs de classe).")]))

    C("« __str__ » et « __repr__ »",
      ("txt", [("p", "Deux méthodes spéciales définissent la représentation en chaîne d'un objet :"),
               ("b0", "`__str__` — la version **lisible**, pour l'utilisateur : appelée par `print(objet)` et `str(...)`"),
               ("b0", "`__repr__` — la version **technique**, pour le développeur : idéalement, de quoi recréer l'objet")]),
      ("code", """
class Voiture:
    def __init__(self, marque):
        self.marque = marque

    def __str__(self):
        return f"Voiture {self.marque}"

    def __repr__(self):
        return f"Voiture(marque={self.marque!r})"
"""))

    C("__str__ vs __repr__ en situation",
      ("txt", [("p", "Le piège classique : `print` d'un objet seul utilise `__str__`... mais `print` d'une **liste** d'objets utilise `__repr__` de chacun !")]),
      ("code", """
v = Voiture("Kia")

print(v)              # Voiture Kia                      ← __str__
print([v, v])         # [Voiture(marque='Kia'), ...]     ← __repr__ !

# Sans ces méthodes, l'affichage par défaut est illisible :
# <__main__.Voiture object at 0x000001E2B3D4F910>
"""),
      ("txt", [("p", "Réflexe professionnel : définir au minimum `__repr__` sur chaque classe métier — le débogage devient immédiatement plus agréable.")]))

    C("Les opérateurs de comparaison",
      ("txt", [("p", "Les méthodes spéciales permettent aussi de redéfinir les opérateurs de comparaison :")]),
      ("table", ["Méthode", "Opérateur", "Méthode", "Opérateur"],
       [["__eq__(self, other)", "==", "__lt__(self, other)", "<"],
        ["__ne__(self, other)", "!=", "__le__(self, other)", "<="],
        ["__gt__(self, other)", ">", "__ge__(self, other)", ">="]],
       [190, 90, 190, 90], {"mono_cols": (0, 1, 2, 3), "row_h": 22}),
      ("gap", 4),
      ("txt", [("p", "Relation implicite : `__ne__` délègue à `__eq__` et renvoie l'inverse — définir `__eq__` suffit pour `==` et `!=`.")]))

    C("« __eq__ » en pratique",
      ("txt", [("p", "Sans `__eq__`, `==` compare **l'identité** (comme `is`) — pas le contenu ! Question d'entretien : « que vaut `==` entre deux objets sans `__eq__` ? »")]),
      ("code", """
class Voiture:
    def __init__(self, plaque):
        self.plaque = plaque

v1 = Voiture("1-ABC-123")
v2 = Voiture("1-ABC-123")
print(v1 == v2)                # False ?! — identité, pas contenu

class Voiture2:
    def __init__(self, plaque):
        self.plaque = plaque
    def __eq__(self, other):                   # on définit l'égalité métier
        return self.plaque == other.plaque

print(Voiture2("1-ABC-123") == Voiture2("1-ABC-123"))    # True
""", {"size": 11}))

    C("__eq__ et __hash__ — objets dans un set", bonus=True, *[
      ("txt", [("p", "Pour vivre dans un `set` ou servir de clé de `dict`, un objet avec `__eq__` doit aussi définir `__hash__` — cohérent avec son égalité :")]),
      ("code", """
class Voiture:
    def __init__(self, plaque):
        self.plaque = plaque
    def __eq__(self, other):
        return self.plaque == other.plaque
    def __hash__(self):
        return hash(self.plaque)       # même plaque → même hash

parc = {Voiture("1-ABC-123"), Voiture("1-ABC-123")}
print(len(parc))                       # 1 — le doublon est éliminé
""")])

    C("Les conteneurs : __len__, __getitem__, __contains__",
      ("txt", [("p", "Les dunders de **conteneur** : votre objet répond à `len(...)`, aux crochets `[...]` et à l'opérateur `in` — comme une vraie liste.")]),
      ("code", """
class Garage:
    def __init__(self, *voitures):
        self.voitures = list(voitures)
    def __len__(self):                 # len(g)
        return len(self.voitures)
    def __getitem__(self, i):          # g[i]  — et l'itération offerte !
        return self.voitures[i]
    def __contains__(self, v):         # "..." in g
        return v in self.voitures

g = Garage("Clio", "Kia", "Tesla")
print(len(g), g[0], "Kia" in g)        # 3 Clio True
for v in g:                            # le for fonctionne via __getitem__
    print(v, end=" ")                  # Clio Kia Tesla
""", {"size": 11.5}))

    C("Surcharge d'opérateurs : __add__", bonus=True, *[
      ("txt", [("p", "Les opérateurs arithmétiques aussi se redéfinissent — `+` appelle `__add__` :")]),
      ("code", """
class Duree:
    def __init__(self, h, m):
        self.h, self.m = h, m
    def __add__(self, other):              # self + other
        total = (self.h + other.h) * 60 + self.m + other.m
        return Duree(total // 60, total % 60)
    def __repr__(self):
        return f"{self.h}h{self.m:02d}"

print(Duree(1, 40) + Duree(2, 35))         # 4h15
"""),
      ("txt", [("p", "Même famille : `__sub__` (-), `__mul__` (*), `__truediv__` (/) — et `__radd__` quand votre objet est à **droite** de l'addition.")])])

    C("Attributs dynamiques : getattr et setattr", bonus=True, *[
      ("txt", [("p", "Lire ou écrire un attribut dont le **nom n'est connu qu'à l'exécution** (une chaîne de caractères) — le quotidien d'un ORM comme celui d'Odoo :")]),
      ("code", """
class Voiture:
    def __init__(self, marque):
        self.marque = marque

v = Voiture("Kia")
print(hasattr(v, "marque"))            # True — l'attribut existe ?
print(getattr(v, "marque"))            # Kia  — lecture par NOM (une chaîne !)
print(getattr(v, "prix", 0))           # 0    — valeur par défaut si absent

setattr(v, "annee", 2024)              # écriture par nom : v.annee = 2024
print(v.annee)                         # 2024

champ = "marque"                       # le nom peut venir d'une variable,
print(getattr(v, champ))               # d'un fichier, d'une base de données...
""", {"size": 11.5})])

    C("@dataclass — l'usine à classes de données", bonus=True, *[
      ("txt", [("p", "Pour une classe qui transporte surtout des données, le décorateur `@dataclass` génère automatiquement `__init__`, `__repr__` et `__eq__` :")]),
      ("code", """
from dataclasses import dataclass

@dataclass
class Voiture:
    marque: str
    annee: int
    prix: float = 0.0          # valeur par défaut possible

v = Voiture("Kia", 2024, 15900.0)
print(v)                       # Voiture(marque='Kia', annee=2024, prix=15900.0)
print(v == Voiture("Kia", 2024, 15900.0))     # True — __eq__ offert !
""")])

    C("__slots__ — des objets au régime", bonus=True, *[
      ("txt", [("p", "Par défaut, chaque objet transporte son `__dict__` — souple, mais coûteux en mémoire. `__slots__` **fige la liste des attributs** : plus de dictionnaire, des objets plus légers et des accès plus rapides.")]),
      ("code", """
class Point:
    __slots__ = ("x", "y")             # la liste FERMÉE des attributs

    def __init__(self, x, y):
        self.x, self.y = x, y

p = Point(3, 4)
print(p.x, p.y)                        # 3 4
# p.z = 5                  # ❌ AttributeError : pas dans __slots__
# print(p.__dict__)        # ❌ AttributeError : plus de __dict__ !
"""),
      ("txt", [("p", "À réserver aux classes instanciées **en masse** (des millions de points, de lignes...) — et toujours mesurer avant d'optimiser.")])])

    C("__call__ et __new__", bonus=True, *[
      ("txt", [("p", "Deux dunders de « culture d'entretien » :"),
               ("b0", "`__call__` — rend l'objet **appelable** comme une fonction")]),
      ("code", """
class Compteur:
    def __init__(self):
        self.n = 0
    def __call__(self):                # l'objet devient APPELABLE
        self.n += 1
        return self.n

c = Compteur()
print(c(), c(), c())                   # 1 2 3 — c() comme une fonction !
"""),
      ("txt", [("p", "`__new__` **crée** l'objet en mémoire, puis `__init__` l'**initialise**. « Que se passe-t-il avant `__init__` ? » est un classique d'entretien — on ne redéfinit `__new__` que pour des cas pointus (singleton, types immuables).")])])

    # ======================================================== HÉRITAGE
    SEC("Héritage")

    C("Concept de l'héritage",
      ("txt", [("p", "Une classe peut **hériter** d'une autre : créer une nouvelle classe à partir d'une base existante."),
               ("b0", "la classe qui hérite : **sous-classe**, classe **enfant**, classe fille ou dérivée"),
               ("b0", "la classe dont on hérite : **super-classe** ou classe **mère**"),
               ("p", "L'enfant récupère **tous** les attributs et méthodes de la mère — puis ajoute ou spécialise. La relation se lit « **est un** » : une voiture électrique **est une** voiture.")]))

    C("L'héritage en Python",
      ("txt", [("p", "L'héritage se déclare sur la classe **enfant**, entre parenthèses. La fonction `issubclass` teste la hiérarchie entre les types.")]),
      ("code", """
class Voiture:
    def __init__(self, marque):
        self.marque = marque

class VoitureElectrique(Voiture):      # hérite de Voiture
    pass

ve = VoitureElectrique("Tesla")
print(ve.marque)                       # Tesla — hérité sans rien écrire !
print(issubclass(VoitureElectrique, Voiture))      # True
"""))

    C("La fonction « super() »",
      ("txt", [("p", "`super()` donne accès aux méthodes de la classe mère depuis l'enfant — notamment pour **compléter** `__init__` au lieu de le réécrire :")]),
      ("code", """
class Voiture:
    def __init__(self, marque):
        self.marque = marque

class VoitureElectrique(Voiture):
    def __init__(self, marque, autonomie):
        super().__init__(marque)       # la mère initialise SA partie...
        self.autonomie = autonomie     # ... l'enfant ajoute la sienne

ve = VoitureElectrique("Tesla", 500)
print(ve.marque, ve.autonomie)         # Tesla 500
"""))

    C("Redéfinition de méthodes",
      ("txt", [("p", "L'enfant peut **redéfinir** une méthode de la mère pour spécialiser le comportement — et la version de l'objet le plus spécifique gagne.")]),
      ("code", """
class Voiture:
    def demarrer(self):
        return "Vroum"

class VoitureElectrique(Voiture):
    def demarrer(self):                        # redéfinition
        return "... (silence électrique)"

print(Voiture().demarrer())                    # Vroum
print(VoitureElectrique().demarrer())          # ... (silence électrique)
"""),
      ("txt", [("p", "Pour **enrichir** plutôt que remplacer : rappeler la version de la mère avec `super().demarrer()` dans la redéfinition.")]))

    C("Le polymorphisme",
      ("txt", [("p", "Le polymorphisme : un objet peut avoir **plusieurs types** — une `VoitureElectrique` est aussi une `Voiture`. La fonction `isinstance` teste les types d'un objet, héritage compris.")]),
      ("code", """
ve = VoitureElectrique("Tesla", 500)

print(isinstance(ve, VoitureElectrique))   # True
print(isinstance(ve, Voiture))             # True — aussi ! c'est le polymorphisme

# Conséquence : un code écrit pour des Voiture accepte
# toutes leurs sous-classes, sans modification.
"""))

    C("L'héritage multiple",
      ("txt", [("p", "Python permet d'hériter de **plusieurs** classes (ce que C# et Java interdisent). Puissant — et source de questions :"),
               ("b0", "si deux mères définissent la même méthode, laquelle est utilisée ?"),
               ("b0", "comment cibler la méthode d'un parent précis ?"),
               ("p", "La réponse de Python : le **MRO** (*Method Resolution Order*) — un ordre de résolution déterministe, consultable via `__mro__`.")]))

    C("MRO : l'ordre des parents compte",
      ("code", """
class Car:
    def description(self):
        return "je roule"

class Plane:
    def description(self):
        return "je vole"

class FlyingCar(Car, Plane):               # Car d'abord !
    pass

print(FlyingCar().description())           # je roule — le PREMIER parent gagne
print([c.__name__ for c in FlyingCar.__mro__])
# ['FlyingCar', 'Car', 'Plane', 'object']
"""),
      ("txt", [("p", "`super()` suit cet ordre. Pour cibler un parent **précis**, on appelle par son nom : `Plane.description(self)` — en passant `self` soi-même. Question d'entretien récurrente.")]))

    C("Héritage multiple : étude de cas vécue",
      ("txt", [("p", "Le cas réel d'une `FlyingCar(Car, Plane)` — chaque parent initialise **sa** partie, par appel explicite :")]),
      ("code", """
class FlyingCar(Car, Plane):
    def __init__(self, name, year, id, nb_motor):
        Car.__init__(self, name, year)              # appel par NOM → self explicite
        Plane.__init__(self, id, year, nb_motor)    # ... sans oublier year !

# Les deux bugs classiques (vécus !) :
# 1) Plane.__init__(self, id, nb_motor)  → year oublié : les positionnels
#    GLISSENT d'un cran et l'erreur accuse... le DERNIER paramètre :
#    TypeError: missing 1 required positional argument: 'nb_motor'
# 2) print(Car.roule())  → sans self :
#    TypeError: roule() missing 1 required positional argument: 'self'
""", {"size": 11.5}),
      ("txt", [("p", "Parade : appeler avec des **arguments nommés** (`year=year`) — un décalage silencieux devient impossible.")]))

    C("Le mixin — l'héritage multiple bien employé",
      ("txt", [("p", "Le **mixin** : une petite classe **sans état propre**, qui n'existe que pour apporter une capacité réutilisable. C'est l'usage sain de l'héritage multiple — Odoo en est truffé.")]),
      ("code", """
class ComparableParPrix:                   # ni __init__ ni attribut :
    def moins_cher_que(self, other):       # il APPORTE une capacité, c'est tout
        return self.prix < other.prix

class Voiture(ComparableParPrix):
    def __init__(self, marque, prix):
        self.marque, self.prix = marque, prix

class Maison(ComparableParPrix):           # la MÊME capacité, greffée ailleurs
    def __init__(self, ville, prix):
        self.ville, self.prix = ville, prix

print(Voiture("Kia", 15900).moins_cher_que(Maison("Mons", 250000)))   # True
""", {"size": 11}),
      ("txt", [("p", "Convention : le suffixe `Mixin`, placé **avant** la classe de base (cf. MRO).")]))

    C("Composition VS héritage",
      ("txt", [("p", "L'héritage n'est pas le seul outil de réutilisation — souvent, la **composition** est le meilleur choix :"),
               ("b0", "héritage — « **est un** » : une voiture électrique est une voiture"),
               ("b0", "composition — « **a un** » : une voiture a un moteur (elle n'est pas un moteur !)")]),
      ("code", """
class Moteur:
    def __init__(self, chevaux):
        self.chevaux = chevaux

class Voiture:                             # Voiture N'hérite PAS de Moteur...
    def __init__(self, marque, chevaux):
        self.marque = marque
        self.moteur = Moteur(chevaux)      # ... elle EN CONTIENT un

print(Voiture("Kia", 136).moteur.chevaux)  # 136
""", {"size": 11.5}),
      ("txt", [("p", "Question d'entretien : « héritage ou composition ? » — répondre par le test « est un / a un ».")]))

    C("L'héritage au quotidien : vos exceptions",
      ("txt", [("p", "Le premier héritage que vous écrirez en entreprise : une **exception métier**, qui hérite d'`Exception` — un nom parlant, un `except` sur mesure.")]),
      ("code", """
class SoldeInsuffisantError(Exception):
    \"\"\"Levée quand un retrait dépasse le solde disponible.\"\"\"

def retirer(solde, montant):
    if montant > solde:
        raise SoldeInsuffisantError(f"Retrait {montant} refusé : solde {solde}")
    return solde - montant

try:
    retirer(50, 200)
except SoldeInsuffisantError as e:         # un except au nom métier
    print(e)
"""))

    # ======================================================== CLASSES ABSTRAITES
    SEC("Les classes abstraites")

    C("L'objectif",
      ("txt", [("p", "Parfois, la classe mère ne peut pas fournir une implémentation complète : elle veut **imposer** une méthode à ses enfants sans l'écrire elle-même."),
               ("b0", "une **méthode abstraite** = la signature dans la mère, l'implémentation déléguée aux enfants"),
               ("b0", "une classe qui possède une méthode abstraite **ne peut pas être instanciée**"),
               ("p", "En Python, c'est le module `abc` (*Abstract Base Class*) qui fournit l'infrastructure.")]))

    C("Le module « abc »",
      ("code", """
from abc import ABC, abstractmethod

class Forme(ABC):                          # hérite d'ABC
    @abstractmethod
    def aire(self): ...                    # signature seule : à implémenter !

class Carre(Forme):
    def __init__(self, cote):
        self.cote = cote
    def aire(self):                        # l'enfant honore le contrat
        return self.cote ** 2

print(Carre(3).aire())     # 9
# Forme()                  # ❌ TypeError : classe abstraite non instanciable
"""),
      ("txt", [("p", "Un enfant qui « oublie » d'implémenter `aire` devient abstrait à son tour — et refuse lui aussi de s'instancier : le contrat est vérifié à la création.")]))

    # ======================================================== MEMBRES STATIQUES
    SEC("Les membres statiques")

    C("L'objectif",
      ("txt", [("p", "Les membres statiques appartiennent **au type lui-même** plutôt qu'à une instance — on y accède par le nom de la classe."),
               ("p", "⚠ Particularité Python : contrairement à C# et Java, un membre statique est **aussi accessible depuis une instance**... et une affectation via l'instance le **masque** (cf. le piège du masquage, section classes)."),
               ("p", "Usage type : un compteur, une constante de classe, un utilitaire lié au type.")]))

    C("Attribut statique : le compteur d'instances",
      ("code", """
class Voiture:
    nb_creees = 0                          # le compteur vit sur la CLASSE

    def __init__(self, marque):
        self.marque = marque
        Voiture.nb_creees += 1             # via la CLASSE — pas self.nb_creees !

Voiture("Renault"); Voiture("Kia"); Voiture("Tesla")
print(Voiture.nb_creees)                   # 3
"""),
      ("txt", [("p", "Pourquoi pas `self.nb_creees += 1` ? Parce que cela créerait un attribut d'instance **masquant** le compteur partagé — chaque voiture compterait... jusqu'à 1.")]))

    C("Méthode statique : @staticmethod",
      ("txt", [("p", "Une méthode statique n'utilise **aucune donnée d'instance** : pas de `self`. Elle vit dans la classe parce qu'elle s'y rattache logiquement.")]),
      ("code", """
class Conversion:
    @staticmethod
    def km_en_miles(km):                   # pas de self !
        return round(km * 0.621371, 1)

print(Conversion.km_en_miles(100))         # 62.1 — appel via la classe
"""))

    C("@classmethod et « cls »",
      ("txt", [("p", "Entre les deux : la **méthode de classe** reçoit la classe elle-même (`cls`). Son grand usage : les **constructeurs alternatifs**.")]),
      ("code", """
class Voiture:
    def __init__(self, marque, annee):
        self.marque = marque
        self.annee = annee

    @classmethod
    def depuis_dict(cls, d):               # cls = la classe (Voiture)
        return cls(d["marque"], d["annee"])

v = Voiture.depuis_dict({"marque": "Kia", "annee": 2024})
print(v.marque, v.annee)                   # Kia 2024
""", {"size": 11.5}))

    C("De l'objet au JSON (et retour)",
      ("txt", [("p", "Le pont entre vos classes et vos fichiers : `__dict__` transforme l'objet en dictionnaire (→ `json.dumps`), et un `@classmethod` refait l'objet depuis le dictionnaire (`cls(**d)`).")]),
      ("code", """
import json

class Voiture:
    def __init__(self, marque, annee):
        self.marque = marque
        self.annee = annee
    @classmethod
    def depuis_dict(cls, d):
        return cls(**d)                        # dict → objet

v = Voiture("Kia", 2024)
texte = json.dumps(v.__dict__)                 # objet → dict → JSON
print(texte)                                   # {"marque": "Kia", "annee": 2024}
v2 = Voiture.depuis_dict(json.loads(texte))    # JSON → dict → objet
print(v2.marque, v2.annee)                     # Kia 2024
""", {"size": 11}))

    C("self, cls... ou rien ? Le récapitulatif",
      ("gap", 6),
      ("table", ["Décorateur", "1er paramètre", "Accède à"],
       [["(aucun)", "self — l'instance", "les attributs de l'objet"],
        ["@classmethod", "cls — la classe", "les attributs de classe, les constructeurs alternatifs"],
        ["@staticmethod", "aucun", "rien : une fonction rangée dans la classe"]],
       [150, 180, 330], {"mono_cols": (0,), "row_h": 26}),
      ("gap", 10),
      ("txt", [("p", "Le bon réflexe pour choisir :"),
               ("b0", "la méthode utilise `self` → méthode d'instance (le cas normal)"),
               ("b0", "elle crée ou décrit **la classe** → `@classmethod`"),
               ("b0", "elle n'utilise ni l'un ni l'autre → `@staticmethod`")]))

    # ======================================================== ÉNUMÉRATIONS
    SEC("Les énumérations")

    C("Définir une énumération",
      ("txt", [("p", "Une énumération fige un **ensemble fermé de valeurs nommées**. Fini les chaînes magiques `\"valide\"` éparpillées dans le code — et leurs fautes de frappe silencieuses :")]),
      ("code", """
from enum import Enum

class Statut(Enum):
    EN_ATTENTE = 1
    VALIDEE = 2
    ANNULEE = 3

print(Statut.VALIDEE)              # Statut.VALIDEE
print(Statut.VALIDEE.name)         # VALIDEE
print(Statut.VALIDEE.value)        # 2
print(Statut(3))                   # Statut.ANNULEE — retrouvé par la valeur
"""),
      ("txt", [("p", "Un membre inexistant lève immédiatement une erreur : rien ne passe inaperçu.")]))

    C("Utiliser une énumération",
      ("txt", [("p", "Comparaison par **identité** (`is`), itération, `auto()` pour la numérotation :")]),
      ("code", """
from enum import Enum, auto

class Direction(Enum):
    NORD = auto()                      # auto() numérote tout seul (1, 2...)
    SUD = auto()
    EST = auto()
    OUEST = auto()

d = Direction.NORD
if d is Direction.NORD:                # comparaison par IDENTITÉ : is
    print("cap au nord")

for d in Direction:                    # une énumération s'itère
    print(d.name, end=" ")             # NORD SUD EST OUEST
""", {"size": 11}),
      ("txt", [("p", "Parfaits pour des états métier — et utilisables en clé de dictionnaire.")]))

    # ======================================================== INTERFACES
    SEC("Les interfaces")

    C("L'objectif",
      ("txt", [("p", "Une interface décrit un **ensemble de fonctionnalités liées** que n'importe quelle classe peut promettre : un « contrat » à respecter."),
               ("p", "Python ne supporte pas nativement les interfaces. Trois approches la remplacent :"),
               ("b0", "le **duck typing** — la philosophie historique de Python"),
               ("b0", "les **classes abstraites** (ABC) — le contrat vérifié à l'instanciation"),
               ("b0", "`typing.Protocol` — le contrat vérifié par l'éditeur (moderne)")]))

    C("Le Duck Typing",
      ("txt", [("p", "Le nom vient du test du canard : « si je vois un oiseau qui vole comme un canard et cancane comme un canard, alors j'appelle cet oiseau un canard »."),
               ("p", "Concrètement : si un objet **définit les bonnes méthodes**, il fait l'affaire — peu importe sa classe.")]),
      ("code", """
class Canard:
    def crier(self):
        return "Coin coin"

class Robot:                       # aucun lien d'héritage avec Canard
    def crier(self):
        return "Bip bip"

for truc in [Canard(), Robot()]:
    print(truc.crier())            # il sait crier : ça suffit !
"""))

    C("Simuler une interface avec ABC",
      ("txt", [("p", "L'approche « contrat strict » : une classe abstraite qui ne contient **que des signatures**, héritée par les classes concrètes."),
               ("b0", "créer l'« interface » : une classe ABC aux méthodes toutes abstraites"),
               ("b0", "en hériter et implémenter les méthodes"),
               ("b0", "vérifier le contrat avec `isinstance(objet, MonInterface)` — grâce au polymorphisme")]))

    C("typing.Protocol — le contrat moderne",
      ("txt", [("p", "`Protocol` formalise le duck typing : on décrit **la forme attendue**, et l'éditeur (VS Code, mypy) vérifie la conformité — **sans aucun héritage**.")]),
      ("code", """
from typing import Protocol

class PeutCrier(Protocol):         # le contrat : juste des signatures
    def crier(self) -> str: ...

def faire_crier(truc: PeutCrier):  # « j'accepte tout ce qui a cette forme »
    print(truc.crier())

faire_crier(Canard())              # accepté SANS héritage : bonne forme !
"""),
      ("txt", [("p", "Le meilleur des deux mondes : la souplesse du duck typing, la vérification d'un contrat. C'est l'approche des bibliothèques Python modernes.")]))

    # ======================================================== FIN
    C("Pour aller plus loin",
      ("txt", [("p", "Ce module couvre l'essentiel de l'objet en Python. Les suites naturelles :"),
               ("b0", "**@dataclass en profondeur** — champs, immutabilité (`frozen=True`), tri automatique"),
               ("b0", "**Les décorateurs** — comprendre enfin ce que font `@property` et ses cousins"),
               ("b0", "**Les tests avec pytest** — tester ses classes méthode par méthode"),
               ("b0", "**Les design patterns** — Factory, Strategy, Observer : l'objet au niveau architecture"),
               ("b0", "**Les ORM** — SQLAlchemy, modèles Odoo : quand vos classes deviennent des tables")]),
      body_size=15.5)

    n += 1
    slide_merci(prs, n)

    prs.save(OUT)
    print(f"OK — {n} slides → {OUT}")


if __name__ == "__main__":
    build()
