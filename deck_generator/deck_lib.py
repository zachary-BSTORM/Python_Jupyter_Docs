"""Moteur de mise en page — style bstorm Training (reproduit du deck Orienté objet)."""
import io
import keyword
import re
import tokenize
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ASSETS = Path(__file__).parent / "assets"

# ---------------------------------------------------------------- constantes
GREEN = RGBColor(0x23, 0xAC, 0x67)
PURPLE = RGBColor(0x3B, 0x2C, 0x87)
TEXT = RGBColor(0x0D, 0x0D, 0x0D)
GRAY = RGBColor(0x8C, 0x8C, 0x8C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
BORDER = "999999"

KW_COLOR = RGBColor(0x00, 0x00, 0xBF)      # mots-clés : bleu foncé gras
STR_COLOR = RGBColor(0x00, 0x83, 0x8A)     # strings : teal gras
COM_COLOR = RGBColor(0x8C, 0x8C, 0x8C)     # commentaires : gris italique
NUM_COLOR = RGBColor(0x16, 0x16, 0xC8)     # nombres : bleu

TITLE_FONT = "Zen Maru Gothic"
BODY_FONT = "Barlow"
CODE_FONT = "Consolas"

SLIDE_W, SLIDE_H = Pt(720), Pt(405)


# ---------------------------------------------------------------- présentation
def new_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


# ---------------------------------------------------------------- primitives
def add_textbox(slide, x, y, w, h):
    box = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return box, tf


def style_run(r, text, font=BODY_FONT, size=16, color=TEXT, bold=False, italic=False):
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.italic = italic
    return r


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def rich_runs(p, text, size=16, color=TEXT, bold=False, font=BODY_FONT):
    """Écrit `text` dans le paragraphe p en interprétant **gras** et `code`."""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        r = p.add_run()
        if part.startswith("**") and part.endswith("**"):
            style_run(r, part[2:-2], font=font, size=size, color=color, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            style_run(r, part[1:-1], font=CODE_FONT, size=size - 1.5, color=color, bold=bold)
        else:
            style_run(r, part, font=font, size=size, color=color, bold=bold)


def set_bullet(p, char="●", level=0, size_pct=100):
    """Transforme le paragraphe en puce avec retrait négatif (vraie puce)."""
    pPr = p._p.get_or_add_pPr()
    if level == 0:
        pPr.set("marL", str(Pt(18)))
        pPr.set("indent", str(-Pt(18)))
    else:
        pPr.set("marL", str(Pt(46)))
        pPr.set("indent", str(-Pt(16)))
    buSzPct = etree.SubElement(pPr, qn("a:buSzPct"))
    buSzPct.set("val", str(size_pct * 1000))
    buFont = etree.SubElement(pPr, qn("a:buFont"))
    buFont.set("typeface", "Arial")
    buChar = etree.SubElement(pPr, qn("a:buChar"))
    buChar.set("char", char)


def no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    etree.SubElement(pPr, qn("a:buNone"))


# ---------------------------------------------------------------- chrome
def add_footer(slide, num):
    """Petit logo bstorm en bas à gauche + numéro de page vert en bas à droite."""
    slide.shapes.add_picture(str(ASSETS / "logo_small.png"), Pt(12), Pt(382), height=Pt(17))
    box, tf = add_textbox(slide, 676, 383, 34, 16)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    style_run(p.add_run(), str(num), font="Arial", size=10, color=GREEN, bold=True)


def add_slide_title(slide, text):
    box, tf = add_textbox(slide, 30, 34, 660, 44)
    p = tf.paragraphs[0]
    rich_runs(p, text, size=26, color=GREEN, font=TITLE_FONT)


def add_bonus_tag(slide):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(608), Pt(16), Pt(82), Pt(21))
    shp.adjustments[0] = 0.5
    shp.fill.solid()
    shp.fill.fore_color.rgb = PURPLE
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    style_run(p.add_run(), "🎓 BONUS", font=BODY_FONT, size=10, color=WHITE, bold=True)


# ---------------------------------------------------------------- code
_KW = set(keyword.kwlist)


def _tokenize_styles(code):
    """Retourne {n° ligne: [(col_debut, col_fin, style), ...]} via tokenize."""
    lines = code.split("\n")
    styles = {i: [] for i in range(len(lines))}
    fstr_types = {getattr(tokenize, n, -1) for n in
                  ("FSTRING_START", "FSTRING_MIDDLE", "FSTRING_END")}
    for tok in tokenize.generate_tokens(io.StringIO(code + "\n").readline):
        ttype, tval, (r1, c1), (r2, c2), _ = tok
        if ttype == tokenize.COMMENT:
            style = "com"
        elif ttype == tokenize.STRING or ttype in fstr_types:
            style = "str"
        elif ttype == tokenize.NUMBER:
            style = "num"
        elif ttype == tokenize.NAME and tval in _KW:
            style = "kw"
        else:
            continue
        for row in range(r1 - 1, r2):
            if row >= len(lines):
                continue
            start = c1 if row == r1 - 1 else 0
            end = c2 if row == r2 - 1 else len(lines[row])
            if start < end:
                styles[row].append((start, end, style))
    return styles


_FALLBACK_KW = re.compile(r"\b(" + "|".join(sorted(_KW, key=len, reverse=True)) + r")\b")


def _fallback_styles(code):
    """Coloration approximative si tokenize échoue (regex, ligne par ligne)."""
    styles = {}
    for i, line in enumerate(code.split("\n")):
        spans = []
        com = line.find("#")
        scan = line if com < 0 else line[:com]
        if com >= 0:
            spans.append((com, len(line), "com"))
        for m in re.finditer(r"(\"[^\"]*\"|'[^']*')", scan):
            spans.append((m.start(), m.end(), "str"))
        for m in _FALLBACK_KW.finditer(scan):
            if not any(s <= m.start() < e for s, e, _ in spans):
                spans.append((m.start(), m.end(), "kw"))
        styles[i] = spans
    return styles


CODE_STYLE = {
    "kw": dict(color=KW_COLOR, bold=True),
    "str": dict(color=STR_COLOR, bold=True),
    "com": dict(color=COM_COLOR, italic=True),
    "num": dict(color=NUM_COLOR),
}


def add_code(slide, code, x, y, w, size=12):
    """Boîte de code façon capture d'écran : blanc, bordure grise, code colorisé."""
    code = code.strip("\n")
    lines = code.split("\n")
    try:
        styles = _tokenize_styles(code)
    except Exception:
        print(f"  [!] tokenize a échoué, fallback regex : {lines[0][:40]}...")
        styles = _fallback_styles(code)

    maxlen = max(len(l) for l in lines) or 1
    # réduire la police si la ligne la plus longue déborde
    fit = (w - 24) / (maxlen * 0.616)
    size = min(size, round(fit * 2) / 2)
    line_h = size * 1.3
    h = len(lines) * line_h + 14

    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = WHITE
    shp.line.color.rgb = RGBColor(0x99, 0x99, 0x99)
    shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Pt(9)
    tf.margin_top = tf.margin_bottom = Pt(6)

    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = Pt(line_h)
        spans = sorted(styles.get(i, []))
        cursor = 0
        for start, end, st in spans:
            if start > cursor:
                style_run(p.add_run(), line[cursor:start], font=CODE_FONT, size=size, color=TEXT)
            kw = CODE_STYLE[st]
            style_run(p.add_run(), line[start:end], font=CODE_FONT, size=size,
                      color=kw.get("color", TEXT), bold=kw.get("bold", False),
                      italic=kw.get("italic", False))
            cursor = end
        if cursor < len(line):
            style_run(p.add_run(), line[cursor:], font=CODE_FONT, size=size, color=TEXT)
        if not line:
            style_run(p.add_run(), "", font=CODE_FONT, size=size, color=TEXT)
    return h


# ---------------------------------------------------------------- tableaux
def _cell_borders(cell, color=BORDER, width_pt=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnB", "a:lnT", "a:lnR", "a:lnL"):     # insérés en tête → ordre L,R,T,B
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = etree.Element(qn(tag))
        ln.set("w", str(Emu(Pt(width_pt))))
        fill = etree.SubElement(ln, qn("a:solidFill"))
        clr = etree.SubElement(fill, qn("a:srgbClr"))
        clr.set("val", color)
        tcPr.insert(0, ln)


def add_table(slide, headers, rows, widths, x, y, size=12, header_size=12,
              row_h=20, header_fill=BLACK, mono_cols=(), align=None):
    """Tableau style bstorm : entête noire/texte blanc, lignes blanches, bordures fines.
    headers=None → pas de ligne d'entête. mono_cols : indices de colonnes en Consolas."""
    n_head = 1 if headers else 0
    n_rows = len(rows) + n_head
    n_cols = len(widths)
    gf = slide.shapes.add_table(n_rows, n_cols, Pt(x), Pt(y), Pt(sum(widths)), Pt(n_rows * row_h))
    tbl = gf.table
    tbl.first_row = False
    tbl.horz_banding = False
    for i, w in enumerate(widths):
        tbl.columns[i].width = Pt(w)
    for r in range(n_rows):
        tbl.rows[r].height = Pt(row_h + (2 if r == 0 and headers else 0))

    def fill_cell(cell, text, *, header=False, mono=False):
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Pt(7)
        cell.margin_right = Pt(5)
        cell.margin_top = cell.margin_bottom = Pt(2)
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_fill if header else WHITE
        _cell_borders(cell)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        if align:
            p.alignment = align
        if header:
            style_run(p.add_run(), text, font=BODY_FONT, size=header_size, color=WHITE, bold=True)
        elif mono:
            style_run(p.add_run(), text, font=CODE_FONT, size=size - 0.5, color=TEXT)
        else:
            rich_runs(p, text, size=size, color=TEXT)

    if headers:
        for c, htext in enumerate(headers):
            fill_cell(tbl.cell(0, c), htext, header=True)
    for r, row in enumerate(rows):
        for c, cell_text in enumerate(row):
            fill_cell(tbl.cell(r + n_head, c), cell_text, mono=c in mono_cols)
    return n_rows * row_h


# ---------------------------------------------------------------- paragraphes / puces
def est_height(text, size, width):
    """Hauteur estimée (pt) d'un paragraphe avec retour à la ligne."""
    plain = re.sub(r"[*`]", "", text)
    cpl = max(10, width / (size * 0.50))
    import math
    return max(1, math.ceil(len(plain) / cpl)) * size * 1.32


def add_paragraphs(slide, items, x, y, w, size=16):
    """items : liste de ('p'|'b0'|'b1', texte). Retourne la hauteur occupée."""
    box, tf = add_textbox(slide, x, y, w, 300)
    total = 0
    for i, (kind, text) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.12
        p.space_after = Pt(6)
        indent_w = w - (18 if kind == "b0" else 46 if kind == "b1" else 0)
        if kind == "b0":
            set_bullet(p, "●", 0, size_pct=70)
        elif kind == "b1":
            set_bullet(p, "○", 1, size_pct=70)
        else:
            no_bullet(p)
        rich_runs(p, text, size=size)
        total += est_height(text, size, indent_w) + 6
    return total
